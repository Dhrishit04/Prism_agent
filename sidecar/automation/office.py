"""Programmatic Microsoft Office document automation with optional COM fallback."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)


class OfficeAutomationError(RuntimeError):
    """Raised when an Office operation cannot be completed."""


class OfficeAutomation:
    """Read and create common Office files without requiring Office to be installed."""

    def __init__(self, use_com: bool = False, allowed_root: str | Path | None = None):
        self.use_com = use_com
        self.allowed_root = Path(allowed_root or Path.cwd()).expanduser().resolve()

    def read_doc(self, path: str) -> dict[str, Any]:
        document_path = self._existing_file(path, {".docx", ".doc"})
        if document_path.suffix.lower() == ".doc" or self.use_com:
            result = self._read_doc_com(document_path)
            if result is not None:
                return result
            if document_path.suffix.lower() == ".doc":
                raise OfficeAutomationError("Reading .doc files requires Microsoft Office and pywin32")

        try:
            from docx import Document
        except ImportError as exc:
            raise OfficeAutomationError("python-docx is required to read .docx files") from exc

        document = Document(str(document_path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        tables = [
            [[cell.text for cell in row.cells] for row in table.rows]
            for table in document.tables
        ]
        return {"path": str(document_path), "paragraphs": paragraphs, "tables": tables}

    def edit_doc(self, path: str, edits: list[dict[str, Any]]) -> dict[str, Any]:
        document_path = self._existing_file(path, {".docx", ".doc"})
        if document_path.suffix.lower() == ".doc" or self.use_com:
            if document_path.suffix.lower() == ".doc" or all(
                edit.get("operation", "replace") in {"replace", "append", "prepend"} for edit in edits
            ):
                result = self._edit_doc_com(document_path, edits)
                if result is not None:
                    return result
            if document_path.suffix.lower() == ".doc":
                raise OfficeAutomationError("Editing .doc files requires Microsoft Office and pywin32")

        try:
            from docx import Document
        except ImportError as exc:
            raise OfficeAutomationError("python-docx is required to edit .docx files") from exc

        document = Document(str(document_path))
        changed = 0
        for edit in edits:
            operation = edit.get("operation", "replace")
            if operation == "replace":
                old = edit.get("old")
                new = edit.get("new", "")
                if not isinstance(old, str) or not old:
                    raise OfficeAutomationError("Replace edits require a non-empty 'old' string")
                for paragraph in document.paragraphs:
                    if old in paragraph.text:
                        self._replace_paragraph_text(paragraph, old, str(new))
                        changed += 1
                for table in document.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            for paragraph in cell.paragraphs:
                                if old in paragraph.text:
                                    self._replace_paragraph_text(paragraph, old, str(new))
                                    changed += 1
            elif operation in {"append", "prepend"}:
                text = edit.get("text")
                if not isinstance(text, str):
                    raise OfficeAutomationError(f"{operation} edits require a 'text' string")
                paragraph = document.add_paragraph()
                if operation == "prepend":
                    document._element.body.insert(0, paragraph._element)
                paragraph.add_run(text)
                changed += 1
            else:
                raise OfficeAutomationError(f"Unsupported document edit operation: {operation}")

        document.save(str(document_path))
        return {"path": str(document_path), "edits_applied": changed}

    def create_pivot(self, csv_path: str, config: dict[str, Any]) -> dict[str, Any]:
        source = self._existing_file(csv_path, {".csv"})
        output_path = self._output_path(config.get("output_path") or source.with_suffix(".pivot.xlsx"))
        try:
            import pandas as pd
        except ImportError as exc:
            raise OfficeAutomationError("pandas and openpyxl are required to create pivot tables") from exc

        index = config.get("index")
        values = config.get("values")
        columns = config.get("columns")
        aggfunc = config.get("aggfunc", "sum")
        if not isinstance(index, (str, list)) or not index:
            raise OfficeAutomationError("Pivot config requires a non-empty 'index'")
        frame = pd.read_csv(source)
        pivot = pd.pivot_table(
            frame,
            index=index,
            values=values,
            columns=columns,
            aggfunc=aggfunc,
            fill_value=config.get("fill_value"),
        )
        output_path.parent.mkdir(parents=True, exist_ok=True)
        pivot.to_excel(output_path, sheet_name=config.get("sheet_name", "Pivot"))
        return {
            "path": str(output_path),
            "rows": int(pivot.shape[0]),
            "columns": int(pivot.shape[1]),
        }

    def read_excel(self, path: str) -> dict[str, Any]:
        workbook_path = self._existing_file(path, {".xlsx", ".xlsm", ".xls"})
        try:
            import pandas as pd
        except ImportError as exc:
            raise OfficeAutomationError("pandas and openpyxl are required to read Excel files") from exc

        sheets = pd.read_excel(workbook_path, sheet_name=None)
        return {
            "path": str(workbook_path),
            "sheets": {
                name: {
                    "columns": list(frame.columns),
                    "rows": frame.astype(object).where(frame.notna(), None).to_dict("records"),
                }
                for name, frame in sheets.items()
            },
        }

    def create_presentation(self, template: str | None, slides: list[dict[str, Any]], output_path: str) -> dict[str, Any]:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise OfficeAutomationError("python-pptx is required to create presentations") from exc

        if template:
            template_path = self._existing_file(template, {".pptx"})
            presentation = Presentation(str(template_path))
        else:
            presentation = Presentation()

        blank_layout = next(
            (layout for layout in presentation.slide_layouts if layout.name.lower() == "blank"),
            presentation.slide_layouts[0],
        )
        for slide_config in slides:
            slide = presentation.slides.add_slide(blank_layout)
            title = slide_config.get("title")
            body = slide_config.get("body", slide_config.get("bullets", []))
            if title:
                title_box = slide.shapes.add_textbox(457200, 228600, 8229600, 685800)
                title_box.text_frame.text = str(title)
            body_box = slide.shapes.add_textbox(457200, 1143000, 8229600, 5029200)
            if isinstance(body, list):
                body_box.text_frame.text = "\n".join(str(item) for item in body)
            else:
                body_box.text_frame.text = str(body)

        destination = self._output_path(output_path)
        destination.parent.mkdir(parents=True, exist_ok=True)
        presentation.save(destination)
        return {"path": str(destination), "slides_created": len(slides)}

    @staticmethod
    def _replace_paragraph_text(paragraph: Any, old: str, new: str) -> None:
        runs = paragraph.runs
        run_text = "".join(run.text for run in runs)
        matches: list[tuple[int, int]] = []
        search_start = 0
        while (match_start := run_text.find(old, search_start)) != -1:
            matches.append((match_start, match_start + len(old)))
            search_start = match_start + len(old)

        offsets: list[tuple[int, int]] = []
        offset = 0
        for run in runs:
            end = offset + len(run.text)
            offsets.append((offset, end))
            offset = end

        for match_start, match_end in reversed(matches):
            start_index = next(
                index for index, (start, end) in enumerate(offsets) if start <= match_start < end
            )
            end_index = next(
                index for index, (start, end) in enumerate(offsets) if start < match_end <= end
            )
            start_offset = match_start - offsets[start_index][0]
            end_offset = match_end - offsets[end_index][0]
            if start_index == end_index:
                runs[start_index].text = (
                    runs[start_index].text[:start_offset]
                    + new
                    + runs[start_index].text[end_offset:]
                )
            else:
                runs[start_index].text = runs[start_index].text[:start_offset] + new
                for index in range(start_index + 1, end_index):
                    runs[index].text = ""
                runs[end_index].text = runs[end_index].text[end_offset:]

    def _existing_file(self, path: str, extensions: set[str]) -> Path:
        if not isinstance(path, str) or not path.strip():
            raise OfficeAutomationError("A file path is required")
        resolved = Path(path).expanduser().resolve()
        self._ensure_allowed(resolved)
        if resolved.suffix.lower() not in extensions:
            raise OfficeAutomationError(f"Unsupported file type: {resolved.suffix or '<none>'}")
        if not resolved.is_file():
            raise OfficeAutomationError(f"File not found: {resolved}")
        return resolved

    def _output_path(self, path: str | Path) -> Path:
        resolved = Path(path).expanduser().resolve()
        self._ensure_allowed(resolved)
        if resolved.suffix.lower() not in {".xlsx", ".pptx"}:
            raise OfficeAutomationError(f"Unsupported output type: {resolved.suffix or '<none>'}")
        return resolved

    def _ensure_allowed(self, path: Path) -> None:
        try:
            path.relative_to(self.allowed_root)
        except ValueError as exc:
            raise OfficeAutomationError(
                f"Access denied: Office path must be inside {self.allowed_root}"
            ) from exc

    def _read_doc_com(self, path: Path) -> dict[str, Any] | None:
        try:
            import win32com.client
        except ImportError:
            return None
        word = win32com.client.Dispatch("Word.Application")
        document = None
        try:
            self._configure_word(word)
            document = word.Documents.Open(str(path), ReadOnly=True, AddToRecentFiles=False)
            paragraphs = [paragraph.Range.Text.rstrip("\r\x07") for paragraph in document.Paragraphs]
            tables = [
                [
                    [cell.Range.Text.rstrip("\r\x07") for cell in row.Cells]
                    for row in table.Rows
                ]
                for table in document.Tables
            ]
            return {"path": str(path), "paragraphs": paragraphs, "tables": tables}
        finally:
            if document is not None:
                document.Close(False)
            word.Quit()

    def _edit_doc_com(self, path: Path, edits: list[dict[str, Any]]) -> dict[str, Any] | None:
        try:
            import win32com.client
        except ImportError:
            return None
        word = win32com.client.Dispatch("Word.Application")
        document = None
        try:
            self._configure_word(word)
            document = word.Documents.Open(str(path), AddToRecentFiles=False)
            changed = 0
            for edit in edits:
                operation = edit.get("operation", "replace")
                if operation == "replace":
                    old = edit.get("old")
                    if not isinstance(old, str) or not old:
                        raise OfficeAutomationError("Replace edits require a non-empty 'old' string")
                    find = document.Content.Find
                    find.Text = old
                    find.Replacement.Text = str(edit.get("new", ""))
                    find.Execute(Replace=2)
                    changed += 1
                elif operation in {"append", "prepend"}:
                    text = edit.get("text")
                    if not isinstance(text, str):
                        raise OfficeAutomationError(f"{operation} edits require a 'text' string")
                    content = document.Content
                    if operation == "append":
                        content.InsertAfter(text)
                    else:
                        content.InsertBefore(text)
                    changed += 1
                else:
                    raise OfficeAutomationError(f"Unsupported document edit operation: {operation}")
            document.Save()
            return {"path": str(path), "edits_applied": changed}
        finally:
            if document is not None:
                document.Close(False)
            word.Quit()

    @staticmethod
    def _configure_word(word: Any) -> None:
        """Open COM documents without running embedded active content."""
        word.Visible = False
        word.DisplayAlerts = 0
        # Office constant msoAutomationSecurityForceDisable.
        word.AutomationSecurity = 3


_office_automation: OfficeAutomation | None = None


def get_office_automation(
    use_com: bool = False, allowed_root: str | Path | None = None
) -> OfficeAutomation:
    global _office_automation
    resolved_root = Path(allowed_root or Path.cwd()).expanduser().resolve()
    if (
        _office_automation is None
        or _office_automation.use_com != use_com
        or _office_automation.allowed_root != resolved_root
    ):
        _office_automation = OfficeAutomation(use_com=use_com, allowed_root=resolved_root)
    return _office_automation
